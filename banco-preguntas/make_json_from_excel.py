# -*- coding: utf-8 -*-
"""Genera questions_1500.json a partir del Excel de Kahoot/Pictionary.
Replica el esquema original y anade los datos de imagen de Kahoot
(estado y URL de origen). Hace backup del JSON anterior si existe.
"""
import json
import os
import shutil

import openpyxl

WORK = os.path.dirname(os.path.abspath(__file__))
os.chdir(WORK)

XLSX = "1500_preguntas_kahoot_pictionary.xlsx"
JSON = "questions_1500.json"
JSON_BACKUP = "questions_1500_backup.json"

if os.path.exists(JSON) and not os.path.exists(JSON_BACKUP):
    shutil.copy2(JSON, JSON_BACKUP)
    print("Backup del JSON anterior:", JSON_BACKUP, flush=True)

wb = openpyxl.load_workbook(XLSX)
ws_k = wb["Kahoot con imagen"]
ws_p = wb["Pictionary"]

n_filas = 1501
preguntas = []
for r in range(2, n_filas + 1):
    qid = ws_k.cell(row=r, column=1).value
    tema = ws_k.cell(row=r, column=2).value
    dificultad = ws_k.cell(row=r, column=3).value
    imagen_pista = ws_k.cell(row=r, column=4).value
    pregunta_k = ws_k.cell(row=r, column=5).value
    opciones = [
        ws_k.cell(row=r, column=6).value,
        ws_k.cell(row=r, column=7).value,
        ws_k.cell(row=r, column=8).value,
        ws_k.cell(row=r, column=9).value,
    ]
    correcta = ws_k.cell(row=r, column=10).value
    nota_imagen = ws_k.cell(row=r, column=11).value
    estado_imagen = ws_k.cell(row=r, column=12).value or ""
    url_imagen = ws_k.cell(row=r, column=14).value or ""

    # Pictionary (misma posicin de fila por ID compartido)
    pid = ws_p.cell(row=r, column=1).value
    if pid != qid:
        raise RuntimeError("IDs desalineados en fila %d: %s vs %s" % (r, qid, pid))
    pista_p = ws_p.cell(row=r, column=4).value
    solucion_p = ws_p.cell(row=r, column=5).value
    ref_p = ws_p.cell(row=r, column=6).value
    regla_p = ws_p.cell(row=r, column=7).value

    preguntas.append({
        "id": qid,
        "tema": tema,
        "dificultad": dificultad,
        "concepto": correcta,
        "respuesta": correcta,
        "kahoot": {
            "pregunta": pregunta_k,
            "imagen_pista": imagen_pista,
            "opciones": opciones,
            "respuesta_correcta": correcta,
            "nota_imagen": nota_imagen,
            "estado_imagen": estado_imagen,
            "imagen_url": url_imagen,
        },
        "pictionary": {
            "pista": pista_p,
            "solucion": solucion_p,
            "pista_referencia": ref_p,
            "regla_sugerida": regla_p,
        },
    })

doc = {
    "titulo": "Banco de 1500 preguntas para Kahoot y Pictionary",
    "version": 2,
    "total_preguntas": len(preguntas),
    "nota": "Generado desde el Excel (1500_preguntas_kahoot_pictionary.xlsx). "
            "La columna imagen_url contiene el origen de la imagen insertada en la hoja 'Kahoot con imagen'; "
            "las descripciones son indicaciones para preparar o buscar imagenes, no archivos incrustados.",
    "questions": preguntas,
}

with open(JSON, "w", encoding="utf-8") as f:
    json.dump(doc, f, ensure_ascii=False, indent=2)

print("LISTO: %s con %d preguntas (%.1f MB)" % (
    JSON, len(preguntas), os.path.getsize(JSON) / 1048576), flush=True)
con_url = sum(1 for p in preguntas if p["kahoot"]["imagen_url"])
print("Con imagen_url: %d / %d" % (con_url, len(preguntas)), flush=True)
print("Ejemplo C0001:", json.dumps(preguntas[0], ensure_ascii=False, indent=1)[:700], flush=True)
