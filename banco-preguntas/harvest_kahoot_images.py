# -*- coding: utf-8 -*-
"""Fase A: descarga la imagen de cada concepto de la hoja 'Kahoot con imagen'.
Fuentes: Wikipedia ES -> Wikipedia EN -> Wikimedia Commons.
Guarda miniaturas JPEG en kahoot_images/ y un cache.json con URL de origen.
Reanudable: los conceptos ya descargados se saltan.
"""
import io
import json
import os
import sys
import time
import urllib.parse
import urllib.request

from PIL import Image

WORK = os.path.dirname(os.path.abspath(__file__))
os.chdir(WORK)

XLSX = os.path.join(WORK, "1500_preguntas_kahoot_pictionary.xlsx")
IMG_DIR = os.path.join(WORK, "kahoot_images")
CACHE_PATH = os.path.join(IMG_DIR, "cache.json")
MAX_BOX = (300, 180)  # tamano maximo de la miniatura guardada (px)

UA = {
    "User-Agent": "KahootPictionaryQuiz/1.0 (use personal; contact: none) urllib",
}


def api_get(url, retries=5):
    for attempt in range(retries):
        try:
            req = urllib.request.Request(url, headers=UA)
            with urllib.request.urlopen(req, timeout=25) as r:
                return json.loads(r.read().decode("utf-8"))
        except urllib.error.HTTPError as e:
            if e.code == 429 and attempt < retries - 1:
                time.sleep(min(25, 5 * (attempt + 1)))
                continue
            print(f"  API fallo: {e}", flush=True)
            return None
        except Exception as e:
            if attempt == retries - 1:
                print(f"  API fallo: {e}", flush=True)
                return None
            time.sleep(1.5 * (attempt + 1))
    return None


def search_wiki(lang, term):
    """Devuelve {'thumb_url','page_title'} del primer resultado con miniatura."""
    url = (
        "https://%s.wikipedia.org/w/api.php?action=query&format=json&formatversion=2"
        "&generator=search&gsrsearch=%s&gsrlimit=4&gsrnamespace=0"
        "&prop=pageimages&piprop=thumbnail&pithumbsize=480"
        % (lang, urllib.parse.quote(term))
    )
    data = api_get(url)
    if not data:
        return None
    pages = data.get("query", {}).get("pages") or []
    pages = [p for p in pages if p.get("thumbnail", {}).get("source")]
    pages.sort(key=lambda p: p.get("index", 99))
    for p in pages:
        th = p["thumbnail"]
        return {
            "thumb_url": th["source"],
            "page_title": p.get("title", ""),
            "lang": lang,
        }
    return None


def search_commons(term):
    url = (
        "https://commons.wikimedia.org/w/api.php?action=query&format=json&formatversion=2"
        "&generator=search&gsrsearch=%s&gsrlimit=4&gsrnamespace=6"
        "&prop=imageinfo&iiprop=url|size&iiurlwidth=480"
        % urllib.parse.quote(term)
    )
    data = api_get(url)
    if not data:
        return None
    pages = data.get("query", {}).get("pages") or []
    pages.sort(key=lambda p: p.get("index", 99))
    for p in pages:
        infos = p.get("imageinfo") or []
        if infos and infos[0].get("thumburl"):
            return {
                "thumb_url": infos[0]["thumburl"],
                "page_title": p.get("title", ""),
                "lang": "commons",
            }
    return None


def download_and_prepare(concept, meta):
    """Descarga la miniatura, la aplana (fondo blanco) y la guarda JPEG."""
    req = urllib.request.Request(meta["thumb_url"], headers=UA)
    with urllib.request.urlopen(req, timeout=30) as r:
        raw = r.read()
    im = Image.open(io.BytesIO(raw))
    im = im.convert("RGBA")
    bg = Image.new("RGB", im.size, (255, 255, 255))
    bg.paste(im, mask=im.split()[3])
    im = bg
    im.thumbnail(MAX_BOX, Image.LANCZOS)
    fname = "%s.jpg" % urllib.parse.quote(concept, safe="")
    fpath = os.path.join(IMG_DIR, fname)
    im.save(fpath, "JPEG", quality=85)
    return {"file": os.path.basename(fpath), "w": im.size[0], "h": im.size[1]}


def load_concepts():
    import openpyxl

    wb = openpyxl.load_workbook(XLSX)
    ws = wb["Kahoot con imagen"]
    concepts = []
    seen = set()
    for r in range(2, ws.max_row + 1):
        c = ws.cell(row=r, column=10).value
        c = (c or "").strip()
        if c and c not in seen:
            seen.add(c)
            concepts.append(c)
    return concepts


def main():
    os.makedirs(IMG_DIR, exist_ok=True)
    cache = {}
    if os.path.exists(CACHE_PATH):
        cache = json.load(open(CACHE_PATH, encoding="utf-8"))

    concepts = load_concepts()
    print("Conceptos unicos: %d" % len(concepts), flush=True)

    with_thumb = 0
    without = []
    for i, concept in enumerate(concepts, 1):
        if concept in cache:
            if cache[concept].get("file"):
                with_thumb += 1
            else:
                without.append(concept)
            continue
        meta = search_wiki("es", concept) or search_wiki("en", concept) or search_commons(concept)
        if meta:
            try:
                info = download_and_prepare(concept, meta)
            except Exception as e:
                print("  descarga fallo '%s': %s" % (concept, e), flush=True)
                info = {}
            if info.get("file"):
                meta["file"] = info["file"]
                meta["w"], meta["h"] = info["w"], info["h"]
                cache[concept] = meta
                with_thumb += 1
            else:
                cache[concept] = {"file": None, "page_title": "", "thumb_url": ""}
                without.append(concept)
        else:
            cache[concept] = {"file": None, "page_title": "", "thumb_url": ""}
            without.append(concept)

        if i % 20 == 0 or i == len(concepts):
            json.dump(cache, open(CACHE_PATH, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
            print("progreso %d/%d | con imagen: %d | sin: %d" % (i, len(concepts), with_thumb, len(without)), flush=True)
        time.sleep(2.0)

    json.dump(cache, open(CACHE_PATH, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("RESUMEN: con imagen %d / sin imagen %d" % (with_thumb, len(without)), flush=True)
    if without:
        print("SIN IMAGEN:", flush=True)
        for c in without:
            print(" -", c, flush=True)


if __name__ == "__main__":
    main()
