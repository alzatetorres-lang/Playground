# -*- coding: utf-8 -*-
"""Fase B: incrusta las imagenes en la hoja 'Kahoot con imagen' del Excel.
- Columna M: imagen incrustada anclada a la fila de la pregunta.
- Columna N: URL de origen de la imagen.
- Columna L: estado ('Insertada' / 'Sin imagen').
Hace backup del Excel original si no existe aun.
"""
import json
import os
import shutil

import openpyxl
from openpyxl.drawing.image import Image as XLImage

WORK = os.path.dirname(os.path.abspath(__file__))
os.chdir(WORK)

XLSX = "1500_preguntas_kahoot_pictionary.xlsx"
BACKUP = "backup_1500_preguntas_kahoot_pictionary.xlsx"
IMG_DIR = "kahoot_images"
CACHE_PATH = os.path.join(IMG_DIR, "cache.json")

IMG_W = 200       # ancho maximo incrustado (px)
IMG_H = 135       # alto maximo incrustado (px)
ROW_PT = 103.0    # alto de fila en puntos (~136 px)

cache = json.load(open(CACHE_PATH, encoding="utf-8"))

if not os.path.exists(BACKUP):
    shutil.copy2(XLSX, BACKUP)
    print("Backup creado:", BACKUP, flush=True)

print("Cargando Excel... (puede tardar un poco)", flush=True)
wb = openpyxl.load_workbook(XLSX)
ws = wb["Kahoot con imagen"]

ws.cell(row=1, column=13, value="Imagen de la respuesta")
ws.cell(row=1, column=14, value="URL origen de la imagen")
ws.column_dimensions["M"].width = 32
ws.column_dimensions["N"].width = 60

filas_con = 0
filas_sin = 0
estilos_der = openpyxl.styles.Alignment(horizontal="left", vertical="center", wrapText=False)

for r in range(2, ws.max_row + 1):
    concept = (ws.cell(row=r, column=10).value or "").strip()
    ws.row_dimensions[r].height = ROW_PT
    meta = cache.get(concept)
    if meta and meta.get("file"):
        fpath = os.path.join(IMG_DIR, meta["file"])
        if os.path.exists(fpath):
            sc = min(IMG_W / meta.get("w", IMG_W), IMG_H / meta.get("h", IMG_H), 1.0)
            img = XLImage(fpath)
            img.width = int((meta.get("w", IMG_W)) * sc)
            img.height = int((meta.get("h", IMG_H)) * sc)
            ws.add_image(img, "M%d" % r)
            ws.cell(row=r, column=13).alignment = estilos_der
            url = meta.get("thumb_url", "")
            ws.cell(row=r, column=14, value=url)
            ws.cell(row=r, column=12, value="Insertada")
            filas_con += 1
            continue
    ws.cell(row=r, column=12, value="Sin imagen")
    filas_sin += 1

print("Escribiendo Excel con %d imagenes... (esto puede tardar)" % filas_con, flush=True)
wb.save(XLSX)
print("LISTO: con imagen %d filas, sin imagen %d filas" % (filas_con, filas_sin), flush=True)
