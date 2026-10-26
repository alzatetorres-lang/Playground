# -*- coding: utf-8 -*-
"""Verificacion final del Excel con imagenes incrustadas."""
import os
import re
import zipfile

WORK = os.path.dirname(os.path.abspath(__file__))
os.chdir(WORK)
X = "1500_preguntas_kahoot_pictionary.xlsx"

z = zipfile.ZipFile(X)
names = z.namelist()
media = set(n for n in names if re.match(r"^xl/media/", n))
drawings = set(n for n in names if re.match(r"^xl/drawings/[^/]+\.xml$", n))

anchors = 0
blips = 0
for d in drawings:
    xml = z.read(d).decode("utf-8")
    anchors += len(re.findall(r"<(?:xdr:)?oneCellAnchor", xml))
    blips += len(re.findall(r'r:embed="(rId\d+)"', xml))


def parse_target(t):
    m = re.match(r"^/xl/media/(.+)$", t, re.I)
    if m:
        return "xl/media/" + m.group(1)
    m = re.match(r"^\.\./media/(.+)$", t, re.I)
    if m:
        return "xl/media/" + m.group(1)
    m = re.match(r"^media/(.+)$", t, re.I)
    if m:
        return "xl/media/" + m.group(1)
    return None


missing = 0
for n in names:
    if re.match(r"^xl/drawings/_rels/", n):
        for m in re.finditer(r'Target="([^"]+)"', z.read(n).decode("utf-8")):
            p = parse_target(m.group(1))
            if p and p not in media:
                missing += 1
print("media: %d | drawings: %d | anchors: %d | blips: %d | targets rotos: %d" % (
    len(media), len(drawings), anchors, blips, missing))

ok_target = missing == 0 and blips == 1500 and anchors >= 1500

import openpyxl
wb = openpyxl.load_workbook(X)
ws = wb["Kahoot con imagen"]
st = {}
sin_url = 0
for r in range(2, ws.max_row + 1):
    v = ws.cell(row=r, column=12).value
    st[v] = st.get(v, 0) + 1
    if v == "Insertada" and not (ws.cell(row=r, column=14).value or "").strip():
        sin_url += 1
print("Columna L (Estado de imagen): %s | Insertadas sin URL: %d" % (st, sin_url))
print("Hojas: %s" % wb.sheetnames)

ct = z.read("[Content_Types].xml").decode("utf-8")
print("Default jpeg presente:", "Extension=\"jpeg\"" in ct or "extension=\"jpeg\"" in ct.lower())
print("Tamano: %.1f MB" % (os.path.getsize(X) / 1048576))
print("VEREDICTO: %s" % ("OK" if (ok_target and st.get("Insertada") == 1500) else "REVISAR"))
